"""Contracts for the real BHUPESH case and future template/proposal extensions."""

from pathlib import Path
from types import SimpleNamespace as Object
from decimal import Decimal
import json, tempfile, unittest, sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "apps/backend"))
from openpyxl import load_workbook
from app.services.passport_service import generate_passport, TEMPLATE_PATH
from app.services.pre_estimate_service import (
    _fonts,
    _zone_pages,
    _gallery,
    _item_detail,
    generate_pre_estimate,
)

ROOT = Path(__file__).resolve().parents[1]


def example():
    source = json.loads((ROOT / "tests/fixtures/bhupesh.json").read_text())
    zones = []
    all_items = []
    keys = [
        "name",
        "description",
        "section",
        "calc_type",
        "quantity",
        "unit",
        "unit_price",
        "total_price",
        "is_tbd",
    ]
    for n, z in enumerate(source["zones"]):
        items = []
        for i, row in enumerate(z["items"]):
            data = dict(zip(keys, row))
            data.update(zone_id=str(n), sort_order=i)
            for key in ("quantity", "unit_price", "total_price"):
                if data[key] is not None:
                    data[key] = Decimal(str(data[key]))
            item = Object(**data)
            items.append(item)
            all_items.append(item)
        zones.append(
            Object(
                label=z["label"],
                sort_order=n,
                line_items=items,
                images=[],
                square_footage=z.get("square_footage"),
                linear_footage=None,
            )
        )
    from datetime import date

    return Object(
        zones=zones,
        line_items=all_items,
        client_name="BHUPESH",
        project_name=None,
        notes=None,
        prepared_by="Kit",
        estimate_date=date(2026, 10, 22),
        type="final",
        title="BHUPESH - Final Estimate",
        final_total=Decimal(89300),
        markup_pct=0,
        discount_pct=0,
    )


class ExportContracts(unittest.TestCase):
    def test_bhupesh_template_and_formulas(self):
        est = example()
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "passport.xlsx"
            generate_passport(est, path)
            ws = load_workbook(path).active
            original = load_workbook(TEMPLATE_PATH).active
            self.assertEqual(ws["C3"].value, "BHUPESH")
            self.assertEqual(ws["C9"].value, 5)
            self.assertIsNone(ws["C2"].value)
            self.assertIsNone(ws["C4"].value)
            expected_rows = [23, 30, 36, 40, 43]
            self.assertEqual(
                ws["I44"].value, "=" + "+".join(f"I{r}" for r in expected_rows)
            )
            self.assertEqual(ws["M5"].value, "=I44")
            self.assertEqual(ws["J58"].value, "=I44/2")
            self.assertEqual(ws["I42"].value, "TBD")
            self.assertIsNone(ws["G42"].value)
            self.assertEqual(ws["C29"].value, "Mirror Delivery & Installation")
            for coordinate in ("F29", "G29", "H29"):
                self.assertIsNone(ws[coordinate].value)
            self.assertIsNone(ws["F39"].value)
            self.assertEqual(ws["G39"].value, 35)
            self.assertEqual(ws["H39"].value, 100)
            self.assertEqual(ws["I39"].value, 3500)
            for cell in ["C20", "D20", "F20", "G20", "H20", "I20"]:
                self.assertEqual(ws[cell]._style, original[cell]._style)
            self.assertEqual(
                set(map(str, ws.merged_cells.ranges)),
                set(map(str, original.merged_cells.ranges)),
            )
            self.assertEqual(ws["O58"].value, original["O58"].value)
            self.assertTrue(ws["M8"].value.startswith("=IFERROR("))
            self.assertTrue(load_workbook(path).calculation.fullCalcOnLoad)

    def test_more_rows_do_not_break_internal_cost_block(self):
        est = example()
        base = est.zones[0].line_items[0]
        est.zones = [
            Object(
                label=f"Zone {n}",
                sort_order=n,
                line_items=[base] * 10,
                images=[],
                square_footage=None,
                linear_footage=None,
            )
            for n in range(10)
        ]
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "large.xlsx"
            generate_passport(est, path)
            ws = load_workbook(path).active
            template = load_workbook(TEMPLATE_PATH).active
            grand = next(row[6].row for row in ws if row[6].value == "TOTAL Estimate")
            self.assertGreater(grand, 110)
            self.assertEqual(ws["M5"].value, f"=I{grand}")
            for row in range(19, 106):
                for col in range(10, 18):
                    expected = template.cell(row, col).value
                    if isinstance(expected, str):
                        expected = expected.replace("I108", f"I{grand}")
                    self.assertEqual(ws.cell(row, col).value, expected)
            self.assertFalse(
                any("#REF!" in str(cell.value) for row in ws for cell in row)
            )

    def test_long_scope_is_paginated_without_shrinking(self):
        _fonts()
        est = example()
        zone = est.zones[0]
        zone.line_items[0].name = "UnbrokenItemName" * 30
        zone.line_items[0].description = (
            "Detailed surface preparation and measured installation scope. " * 150
        )
        pages = _zone_pages(zone, Decimal(1))
        self.assertGreater(len(pages), 3)
        for page in pages:
            capacity = 480 - (len(page.title_lines) - 1) * 36 - 143
            self.assertLessEqual(sum(row.height for row in page.rows), capacity)
            self.assertTrue(
                all(line[2] >= 11 for row in page.rows for line in row.lines)
            )
        self.assertEqual(
            sum(
                bool(row.price)
                for page in pages
                for row in page.rows
                if row.kind == "item"
            ),
            3,
        )
        self.assertEqual(
            " ".join(line[0] for p in pages for r in p.rows for line in r.lines).count(
                "Detailed surface"
            ),
            150,
        )

    def test_images_are_whole_and_keep_aspect_ratio(self):
        from PIL import Image

        path = str(ROOT / "tests/fixtures/bhupesh-reference.png")
        with Image.open(path) as im:
            ratio = im.width / im.height

        class RecordingCanvas:
            def __init__(self):
                self.images = []

            def setFillColor(self, *a):
                pass

            def setStrokeColor(self, *a):
                pass

            def setLineWidth(self, *a):
                pass

            def rect(self, *a, **k):
                pass

            def drawImage(self, image, x, y, w, h, **kwargs):
                self.images.append((image.getSize(), x, y, w, h))

        for count in range(1, 5):
            c = RecordingCanvas()
            _gallery(c, [path] * count, 56, 157, 440, 323)
            self.assertEqual(len(c.images), count)
            for size, x, y, w, h in c.images:
                self.assertAlmostEqual(w / h, ratio)
                self.assertEqual(size, (829, 372))
                self.assertGreaterEqual(x, 56)
                self.assertGreaterEqual(y, 157)
                self.assertLessEqual(x + w, 496)
                self.assertLessEqual(y + h, 480)

    def test_no_photo_export_is_available_and_does_not_mutate_final(self):
        est = example()
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "no-photos.pdf"
            generate_pre_estimate(est, path)
            self.assertTrue(path.read_bytes().startswith(b"%PDF"))
            self.assertEqual(est.type, "final")
            self.assertEqual(est.final_total, 89300)
        item = est.zones[0].line_items[0]
        self.assertIn("$1,045", _item_detail(item, Decimal("1.045")))


if __name__ == "__main__":
    unittest.main()
