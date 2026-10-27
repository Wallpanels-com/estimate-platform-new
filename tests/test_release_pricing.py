from pathlib import Path
import sys, unittest, tempfile
from types import SimpleNamespace as Object
from decimal import Decimal
ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT/'apps/backend'), str(ROOT/'scripts')]
from app.services.calculator_registry import calculate
from app.services.pricing_config import DEFAULT_PRICING_CONFIG
from app.services.passport_service import generate_passport
from app.services.final_pricing import quote
from check_passport import cell_value
from openpyxl import load_workbook
from unittest.mock import AsyncMock, Mock, patch


class ReleasePricing(unittest.TestCase):
    def test_mirror_installation_survives_material_quantity_edit(self):
        result = calculate('mirror', DEFAULT_PRICING_CONFIG, square_footage=20)
        self.assertEqual(result['total'], 2600)
        material, installation = result['line_items']
        self.assertEqual(material['total_price'], 1600)
        self.assertEqual(installation['total_price'], 1000)
        material['quantity'] = 25
        material['total_price'] = material['quantity'] * material['unit_price']
        self.assertEqual(sum(i['total_price'] for i in result['line_items']), 3000)
        custom = {**DEFAULT_PRICING_CONFIG, 'mirror': {**DEFAULT_PRICING_CONFIG['mirror'], 'minimum': 450}}
        self.assertEqual(calculate('mirror', custom, square_footage=20)['total'], 2050)
        self.assertEqual(calculate('mirror', custom, square_footage=20)['line_items'][1]['total_price'], 450)

    def test_excel_matches_final_with_markup_discount_and_manual_total(self):
        for markup, discount, manual in [('0','0',False), ('12.5','5',False), ('0','15',False), ('7.25','3.5',True), ('0','0',True)]:
            with self.subTest(markup=markup, discount=discount, manual=manual), tempfile.TemporaryDirectory() as directory:
                item = Object(id='item', zone_id='zone', name='Panel', description='Finish', quantity=Decimal('3'), unit_price=Decimal('33.33'), total_price=Decimal('99.99'), unit='SF', calc_type='sf', sort_order=0, is_tbd=False)
                zone = Object(id='zone', label='Dining', square_footage=3, linear_footage=None, line_items=[item])
                est = Object(client_name='TEST', project_name='TEST', notes=None, prepared_by='TEST', zones=[zone], line_items=[item], markup_pct=Decimal(markup), discount_pct=Decimal(discount), manual_override=manual, final_total=Decimal('125.50' if manual else '99.99'))
                path = Path(directory)/'passport.xlsx'; generate_passport(est,path)
                sheet = load_workbook(path).active
                grand = next(row[8].coordinate for row in sheet if row[6].value == 'TOTAL Estimate')
                self.assertEqual(cell_value(sheet,grand), quote(est).grand_total)
                self.assertEqual(sheet['M5'].value, '='+grand)


class QuickSavedDocument(unittest.IsolatedAsyncioTestCase):
    async def test_quick_without_photos_is_saved_as_preliminary_pdf(self):
        import uuid
        from app.api.v1 import governance
        with tempfile.TemporaryDirectory() as directory:
            est=Object(id=uuid.uuid4(),type='quick',deleted_at=None,source_estimate_id=None)
            db=Object(execute=AsyncMock(return_value=Object(scalar=lambda:0)),add=Mock(),flush=AsyncMock(),commit=AsyncMock())
            def render(estimate,path):
                Path(path).write_bytes(b'%PDF-quick-preliminary-test')
            with patch.object(governance,'load_estimate',AsyncMock(return_value=est)), \
                 patch.object(governance.settings,'STORAGE_PATH',directory), \
                 patch.object(governance.EstimateOut,'model_validate',return_value=Object(model_dump=lambda **kw:{})), \
                 patch.object(governance,'generate_pre_estimate',side_effect=render) as pre, \
                 patch.object(governance,'_render_final',AsyncMock(side_effect=AssertionError('Quick must not require Final photos'))):
                result=await governance.issue_document(est.id,governance.IssueDocument(kind='pdf',label='TEST quick'),db,Object(name='Employee',email='employee'))
                self.assertEqual(result['kind'],'pdf')
                pre.assert_called_once()
