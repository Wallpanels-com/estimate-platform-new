"""Read-only local BHUPESH verification and API exports; no estimates are created."""

from decimal import Decimal
from pathlib import Path
import hashlib
import json
import httpx
from openpyxl import load_workbook
from qa_credentials import qa_credentials

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / ".local/qa/bhupesh"


def verify():
    OUT.mkdir(parents=True, exist_ok=True)
    eid = (ROOT / ".local/qa/bhupesh-id.txt").read_text().strip()
    client = httpx.Client(
        base_url="http://127.0.0.1:8010/api/v1", trust_env=False, timeout=60
    )
    email, password = qa_credentials()
    response = client.post(
        "/auth/login",
        json={"email": email, "password": password},
    )
    response.raise_for_status()
    client.headers["Authorization"] = "Bearer " + response.json()["access_token"]

    def call(method, path):
        response = client.request(method, path)
        response.raise_for_status()
        return response

    before = call("GET", f"/estimates/{eid}").json()
    fixture = json.loads((ROOT / "tests/fixtures/bhupesh.json").read_text())
    for key, value in fixture["estimate"].items():
        assert before[key] == value, key
    assert (
        Decimal(before["subtotal"])
        == Decimal(before["final_total"])
        == Decimal("89300")
    )
    assert Decimal(before["markup_pct"]) == Decimal(before["discount_pct"]) == 0
    assert len(before["zones"]) == 5
    source_hash = hashlib.sha256(
        (ROOT / "tests/fixtures/bhupesh-reference.png").read_bytes()
    ).hexdigest()
    for zone, source in zip(before["zones"], fixture["zones"]):
        assert zone["label"] == source["label"] and Decimal(
            zone["subtotal"]
        ) == Decimal(source["expected_total"])
        assert (
            len(zone["line_items"]) == len(source["items"]) and len(zone["images"]) == 1
        )
        assert (
            hashlib.sha256(
                Path(zone["images"][0]["file_path"]).read_bytes()
            ).hexdigest()
            == source_hash
        )
        for item, fields in zip(zone["line_items"], source["items"]):
            for key, value in zip(
                [
                    "name",
                    "description",
                    "section",
                    "calc_type",
                    "quantity",
                    "unit",
                    "unit_price",
                    "total_price",
                    "is_tbd",
                ],
                fields,
            ):
                actual = item[key]
                if (
                    key in ("quantity", "unit_price", "total_price")
                    and value is not None
                ):
                    assert Decimal(actual) == Decimal(str(value)), (
                        zone["label"],
                        item["name"],
                        key,
                    )
                else:
                    assert actual == value, (zone["label"], item["name"], key)
    active = call("GET", "/estimates").json()
    assert [e["id"] for e in active] == [eid]
    prices = call("GET", "/pricing/config").json()
    headers = {}
    for suffix, filename in [
        ("pre-estimate", "Pre-Estimate - Wallpanels - BHUPESH.pdf"),
        ("passport", "Project passport - Wallpanels - BHUPESH.xlsx"),
    ]:
        response = call("GET", f"/pdf/{eid}/{suffix}")
        (OUT / filename).write_bytes(response.content)
        headers[suffix] = response.headers.get("content-disposition")
        assert (
            "5.%20Project" not in headers[suffix]
            and "5. Project" not in headers[suffix]
        )
    call("POST", f"/pdf/{eid}/generate")
    (OUT / "legacy-final.pdf").write_bytes(call("GET", f"/pdf/{eid}/download").content)
    after = call("GET", f"/estimates/{eid}").json()
    assert before == after, "Exports changed the saved estimate."
    assert prices == call("GET", "/pricing/config").json(), "Price Book changed."
    workbook = load_workbook(OUT / "Project passport - Wallpanels - BHUPESH.xlsx")
    ws = workbook.active
    assert ws["I44"].value == "=I23+I30+I36+I40+I43"
    assert (
        ws["I42"].value == "TBD" and ws["G42"].value is None and ws["F39"].value is None
    )
    assert not any("#REF!" in str(cell.value) for row in ws for cell in row)
    (OUT / "latest-estimate.json").write_text(json.dumps(after, indent=2))
    result = {
        "id": eid,
        "zones": 5,
        "items": 15,
        "zone_totals": [z["subtotal"] for z in after["zones"]],
        "total": after["subtotal"],
        "single_active_estimate": True,
        "all_fields_match": True,
        "shared_reference_sha256": source_hash,
        "estimate_unchanged_by_exports": True,
        "price_book_unchanged": True,
        "filenames": headers,
    }
    (OUT / "api-verification.json").write_text(json.dumps(result, indent=2))
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    verify()
